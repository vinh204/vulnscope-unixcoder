from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Bao_cao_tien_do_VulnScope.docx"

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"
TEXT = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="8"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def format_run(run, size=12, bold=False, color=TEXT, italic=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def set_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, line=1.3):
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    set_paragraph(p)
    if bold_lead and text.startswith(bold_lead):
        format_run(p.add_run(bold_lead), bold=True)
        format_run(p.add_run(text[len(bold_lead):]))
    else:
        format_run(p.add_run(text))
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_paragraph(p, after=3, line=1.2)
    format_run(p.add_run(text))
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_paragraph(p, after=3, line=1.2)
    format_run(p.add_run(text))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_paragraph(p, WD_ALIGN_PARAGRAPH.LEFT, before=10 if level == 1 else 7, after=6, line=1.15)
    size = 15 if level == 1 else 13
    for run in p.runs:
        format_run(run, size=size, bold=True)
    p.paragraph_format.keep_with_next = True
    return p


def add_table(doc, headers, rows, widths=None, font_size=10.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            cell.width = Inches(widths[i])
        p = cell.paragraphs[0]
        set_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.05)
        format_run(p.add_run(header), size=font_size, bold=True, color=RGBColor(255, 255, 255))
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = Inches(widths[i])
            p = cell.paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or len(str(value)) < 12 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph(p, align, after=0, line=1.08)
            format_run(p.add_run(str(value)), size=font_size)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_pr.extend([r_fonts, color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([r_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    format_run(run, size=10)


def page_break(doc):
    doc.add_page_break()


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(0.85)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.color.rgb = TEXT
    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    # Cover
    for line, size, bold in [
        ("BAN CƠ YẾU CHÍNH PHỦ", 13, True),
        ("HỌC VIỆN KỸ THUẬT MẬT MÃ", 13, True),
    ]:
        p = doc.add_paragraph()
        set_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, after=2, line=1.0)
        format_run(p.add_run(line), size=size, bold=bold)
    doc.add_paragraph().paragraph_format.space_after = Pt(50)
    p = doc.add_paragraph(style="Title")
    set_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, after=18, line=1.15)
    format_run(p.add_run("BÁO CÁO TIẾN ĐỘ ĐỒ ÁN TỐT NGHIỆP"), size=20, bold=True)
    p = doc.add_paragraph()
    set_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, after=12, line=1.2)
    format_run(p.add_run("NGHIÊN CỨU ỨNG DỤNG MÔ HÌNH UNIXCODER\nTRONG PHÁT HIỆN LỖI MÃ NGUỒN"), size=16, bold=True)
    p = doc.add_paragraph()
    set_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, after=32, line=1.1)
    format_run(p.add_run("Sản phẩm demo VulnScope"), size=14, bold=True)
    info = add_table(doc, ["Thông tin", "Nội dung"], [
        ("Sinh viên", "Nguyễn Văn Vinh"),
        ("Mã số sinh viên", "AT19759"),
        ("Lớp", "AT19CP"),
        ("Ngành", "An toàn thông tin"),
        ("Giảng viên hướng dẫn", "TS. Đặng Xuân Bảo"),
    ], widths=[2.0, 4.4])
    doc.add_paragraph().paragraph_format.space_after = Pt(46)
    p = doc.add_paragraph()
    set_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.0)
    format_run(p.add_run("Hà Nội, 2026"), size=12)

    # Overview
    page_break(doc)
    add_heading(doc, "1 Mục đích và phạm vi báo cáo", 1)
    add_body(doc, "Báo cáo này trình bày phần công việc đã hoàn thành sau Chương 1 và Chương 2 của đồ án, tập trung vào sản phẩm demo VulnScope. Nội dung mô tả quy trình xây dựng, các chức năng có thể trình diễn, kết quả kiểm thử, kết quả thực nghiệm ban đầu, hạn chế hiện tại và kế hoạch hoàn thiện trong thời gian tiếp theo.")
    add_body(doc, "Kết luận tiến độ: pipeline từ dữ liệu, checkpoint, suy luận đến giao diện và triển khai trực tuyến đã hoạt động. Demo đã sử dụng checkpoint UniXcoder thật, có thể phân tích mã C/C++, so sánh trước và sau bản vá, giải thích token và trình bày kết quả đánh giá. Tuy nhiên, checkpoint hiện tại mới là baseline kỹ thuật; chất lượng mô hình chưa phải kết quả cuối của đồ án.", bold_lead="Kết luận tiến độ:")

    add_heading(doc, "2 Khối lượng công việc đã hoàn thành", 1)
    rows = [
        ("1", "Chuẩn bị dữ liệu", "Đọc được JSON và JSONL; chuẩn hóa các trường idx, func, target; chuẩn bị split Devign.", "Hoàn thành"),
        ("2", "Huấn luyện", "Fine-tune mô hình phân loại nhị phân, lưu checkpoint tốt nhất theo F1.", "Hoàn thành bản đầu"),
        ("3", "Suy luận", "Nạp checkpoint, dự đoán một mẫu hoặc theo batch, trả về xác suất và độ trễ.", "Hoàn thành"),
        ("4", "Giải thích", "Đánh giá ảnh hưởng token bằng phương pháp token occlusion.", "Hoàn thành"),
        ("5", "Đánh giá", "Tính Accuracy, Precision, Recall, F1, ROC-AUC, PR-AUC và confusion matrix.", "Hoàn thành bản đầu"),
        ("6", "Giao diện", "Xây dựng trang phân tích, so sánh bản vá và dashboard thực nghiệm.", "Hoàn thành"),
        ("7", "Triển khai", "Đưa mã nguồn lên GitHub, checkpoint lên Hugging Face và demo lên Streamlit Cloud.", "Hoàn thành"),
        ("8", "Kiểm thử", "Unit test dữ liệu, metric, bootstrap; smoke test Docker và checkpoint thật.", "Hoàn thành"),
    ]
    add_table(doc, ["STT", "Hạng mục", "Kết quả", "Trạng thái"], rows, widths=[0.45, 1.35, 3.95, 1.05])

    # Architecture
    page_break(doc)
    add_heading(doc, "3 Kiến trúc và quy trình xử lý", 1)
    add_body(doc, "VulnScope được tổ chức theo các mô-đun độc lập để tách quá trình chuẩn bị dữ liệu, huấn luyện, đánh giá và triển khai. Cách tổ chức này giúp kiểm thử từng thành phần và hạn chế việc giao diện can thiệp vào kết quả mô hình.")
    add_heading(doc, "3.1 Các thành phần chính", 2)
    add_table(doc, ["Thành phần", "Vai trò"], [
        ("data_utils.py", "Đọc và chuẩn hóa dữ liệu đầu vào từ JSON hoặc JSONL."),
        ("model.py", "Bao đóng encoder UniXcoder và lớp phân loại nhị phân."),
        ("run.py", "Huấn luyện, đánh giá theo epoch và lưu checkpoint tốt nhất."),
        ("inference.py", "Nạp checkpoint, dự đoán và giải thích token."),
        ("evaluate.py", "Đánh giá test set và tạo artifact cho dashboard."),
        ("app.py", "Giao diện Streamlit cho người dùng cuối."),
        ("bootstrap.py", "Chuẩn bị checkpoint và khởi động ứng dụng production."),
    ], widths=[1.7, 5.1])
    add_heading(doc, "3.2 Luồng xử lý một yêu cầu", 2)
    for text in [
        "Người dùng chọn ví dụ, nhập trực tiếp một hàm hoặc tải file C/C++.",
        "Hệ thống chuẩn hóa khoảng trắng và tokenize bằng tokenizer của UniXcoder.",
        "Chuỗi token được thêm các token điều khiển, cắt theo block size và padding về cùng độ dài.",
        "Mô hình sinh xác suất cho hai lớp safe và vulnerable.",
        "Giao diện so sánh xác suất với ngưỡng cảnh báo và hiển thị kết quả.",
        "Khi người dùng yêu cầu giải thích, hệ thống lần lượt che token để đo thay đổi xác suất.",
    ]:
        add_number(doc, text)

    # Features
    page_break(doc)
    add_heading(doc, "4 Các chức năng của demo VulnScope", 1)
    add_heading(doc, "4.1 Phân tích mã nguồn", 2)
    add_body(doc, "Trang Phân tích tiếp nhận một hàm C/C++ từ ô nhập hoặc file có phần mở rộng c, cc, cpp, h và hpp. Người dùng cũng có thể chọn các ví dụ buffer overflow, use-after-free hoặc phiên bản an toàn. Sau khi phân tích, hệ thống hiển thị xác suất lỗ hổng, độ tin cậy, số token, độ trễ và cảnh báo khi đầu vào bị cắt.")
    for item in [
        "Kết quả được lưu trong session để không biến mất khi giao diện chạy lại.",
        "Người dùng có thể tải kết quả ở định dạng JSON.",
        "Production bắt buộc dùng checkpoint thật và không tự chuyển sang chế độ minh họa.",
        "Kích thước file tải lên được giới hạn để giảm rủi ro tài nguyên.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "4.2 Giải thích token", 2)
    add_body(doc, "Chức năng giải thích sử dụng token occlusion. Hệ thống dự đoán xác suất ban đầu, sau đó lần lượt thay từng token bằng mask token và dự đoán lại. Chênh lệch xác suất được dùng làm attribution. Attribution dương cho biết token làm tăng dự đoán lỗ hổng; giá trị tuyệt đối lớn cho biết token có ảnh hưởng mạnh.")
    add_heading(doc, "4.3 So sánh trước và sau bản vá", 2)
    add_body(doc, "Người dùng nhập hai phiên bản của cùng một hàm. Hệ thống sử dụng cùng checkpoint, block size và ngưỡng để đánh giá hai phiên bản, sau đó hiển thị xác suất trước vá, sau vá, mức thay đổi và phần diff của mã nguồn. Chức năng này hỗ trợ quan sát liệu bản vá có làm giảm rủi ro theo mô hình hay không.")
    add_heading(doc, "4.4 Dashboard thực nghiệm", 2)
    add_body(doc, "Dashboard chỉ đọc artifact do chương trình đánh giá tạo ra. Các chỉ số không được sinh giả trong giao diện. Trang hiện hiển thị Accuracy, Precision, Recall, F1, ROC-AUC, PR-AUC, confusion matrix và biểu đồ phân bố true negative, false positive, false negative, true positive.")

    # Deployment and links
    page_break(doc)
    add_heading(doc, "5 Triển khai và kiểm thử", 1)
    add_heading(doc, "5.1 Triển khai trực tuyến", 2)
    add_body(doc, "Mã nguồn được quản lý trong một repository độc lập trên GitHub. Checkpoint dung lượng khoảng 480 MB được lưu trong Hugging Face Model Repository thay vì commit trực tiếp vào Git. Ứng dụng được triển khai trên Streamlit Community Cloud và tự tải checkpoint đã ghim phiên bản khi khởi động.")
    p = doc.add_paragraph()
    set_paragraph(p, WD_ALIGN_PARAGRAPH.LEFT, after=4)
    format_run(p.add_run("Demo trực tuyến: "), bold=True)
    add_hyperlink(p, "https://vulnscope-unixcoder.streamlit.app", "https://vulnscope-unixcoder.streamlit.app")
    p = doc.add_paragraph()
    set_paragraph(p, WD_ALIGN_PARAGRAPH.LEFT, after=4)
    format_run(p.add_run("Mã nguồn: "), bold=True)
    add_hyperlink(p, "https://github.com/vinh204/vulnscope-unixcoder", "https://github.com/vinh204/vulnscope-unixcoder")
    p = doc.add_paragraph()
    set_paragraph(p, WD_ALIGN_PARAGRAPH.LEFT, after=8)
    format_run(p.add_run("Checkpoint: "), bold=True)
    add_hyperlink(p, "https://huggingface.co/vinh204/vulnscope-unixcoder-devign", "https://huggingface.co/vinh204/vulnscope-unixcoder-devign")
    add_heading(doc, "5.2 Các biện pháp ổn định và an toàn", 2)
    for item in [
        "Đóng gói bằng Docker và chạy container bằng user không có quyền root.",
        "Cung cấp health check để nền tảng giám sát trạng thái ứng dụng.",
        "Dùng cache để chỉ nạp tokenizer và model một lần trong mỗi tiến trình.",
        "Tách checkpoint và dataset lớn khỏi Git repository bằng gitignore và dockerignore.",
        "Không thực thi mã nguồn do người dùng tải lên; hệ thống chỉ tokenize và suy luận.",
        "Ghim revision của checkpoint để deployment có khả năng tái lập.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "5.3 Kết quả kiểm thử", 2)
    add_table(doc, ["Nhóm kiểm thử", "Nội dung", "Kết quả"], [
        ("Unit test", "Đọc JSON, JSONL; metric khi đủ lớp và thiếu lớp; bootstrap checkpoint.", "6 trên 6 đạt"),
        ("Syntax", "Biên dịch kiểm tra toàn bộ thư mục code.", "Đạt"),
        ("Docker", "Build image, khởi động container, kiểm tra health endpoint.", "Đạt"),
        ("Inference", "Nạp checkpoint thật và chạy dự đoán mẫu buffer overflow.", "Đạt"),
        ("Production", "URL public trả HTTP 200 và giao diện hoạt động.", "Đạt"),
    ], widths=[1.3, 4.55, 1.0])

    # Metrics
    page_break(doc)
    add_heading(doc, "6 Kết quả thực nghiệm ban đầu", 1)
    add_body(doc, "Checkpoint hiện tại được xây dựng để xác minh toàn bộ pipeline và phục vụ demo. Encoder được đóng băng, mô hình được huấn luyện 3 epoch trên 2.000 mẫu, block size 128 và đánh giá trên 500 mẫu test. Ngưỡng cảnh báo đang sử dụng là 0,44.")
    add_table(doc, ["Chỉ số", "Giá trị", "Ý nghĩa"], [
        ("Accuracy", "0,504", "Tỷ lệ dự đoán đúng trên toàn bộ mẫu."),
        ("Precision", "0,486", "Tỷ lệ cảnh báo vulnerable thực sự là vulnerable."),
        ("Recall", "0,579", "Tỷ lệ mẫu vulnerable được phát hiện."),
        ("F1", "0,529", "Trung bình điều hòa giữa precision và recall."),
        ("ROC-AUC", "0,528", "Khả năng phân biệt hai lớp trên nhiều ngưỡng."),
        ("PR-AUC", "0,529", "Quan hệ precision-recall trên nhiều ngưỡng."),
    ], widths=[1.35, 1.0, 4.5])
    add_heading(doc, "6.1 Ma trận nhầm lẫn", 2)
    add_table(doc, ["Nhãn thật", "Dự đoán an toàn", "Dự đoán vulnerable"], [
        ("An toàn", "113", "147"),
        ("Vulnerable", "101", "139"),
    ], widths=[2.2, 2.2, 2.45])
    add_body(doc, "Nhận xét: kết quả hiện tại gần mức baseline và chưa đủ để kết luận UniXcoder đạt hiệu quả tốt trên Devign. Kết quả này phản ánh đúng cấu hình thử nghiệm hạn chế: tập huấn luyện nhỏ, encoder đóng băng và block size ngắn. Giá trị của giai đoạn hiện tại là chứng minh pipeline hoạt động từ đầu đến cuối; kết quả cuối cần được xác lập sau khi mở rộng huấn luyện và thực hiện đối chứng.", bold_lead="Nhận xét:")

    # Limitations and plan
    page_break(doc)
    add_heading(doc, "7 Hạn chế và kế hoạch hoàn thiện", 1)
    add_heading(doc, "7.1 Hạn chế hiện tại", 2)
    for item in [
        "Checkpoint chỉ được huấn luyện trên tập con 2.000 mẫu và chưa fine-tune toàn bộ encoder.",
        "Block size 128 có thể làm mất thông tin ở các hàm dài.",
        "Chưa chạy nhiều random seed nên chưa đánh giá được độ ổn định.",
        "Chưa có baseline hoàn chỉnh để so sánh trực tiếp trong cùng điều kiện.",
        "Chưa kiểm tra đầy đủ duplicate và near-duplicate giữa train, validation và test.",
        "Token occlusion có chi phí suy luận cao và chỉ phản ánh mức ảnh hưởng, không chứng minh quan hệ nhân quả.",
        "GitHub Actions chưa chạy được do trạng thái billing của tài khoản; kiểm thử hiện được thực hiện local.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "7.2 Kế hoạch tám tuần tiếp theo", 2)
    add_table(doc, ["Thời gian", "Công việc", "Sản phẩm dự kiến"], [
        ("Tuần 1", "Kiểm tra dữ liệu, duplicate và near-duplicate; chuẩn hóa cấu hình.", "Báo cáo chất lượng split và cấu hình tái lập."),
        ("Tuần 2 đến 3", "Xây dựng TF-IDF + Logistic Regression; mở rộng fine-tune UniXcoder.", "Bảng kết quả baseline và mô hình chính."),
        ("Tuần 4", "Ablation block size, frozen encoder và fine-tune layer cuối.", "Bảng ablation và nhận xét."),
        ("Tuần 5", "Phân tích false positive, false negative; tối ưu threshold.", "Nhóm lỗi điển hình và threshold hợp lý."),
        ("Tuần 6", "Cập nhật checkpoint, dashboard và kiểm thử deployment.", "Demo ổn định với mô hình cải tiến."),
        ("Tuần 7", "Hoàn thiện Chương 3, Chương 4 và kết luận.", "Bản thảo báo cáo đầy đủ."),
        ("Tuần 8", "Đóng băng hệ thống, làm slide và diễn tập.", "Bản nộp, slide và video dự phòng."),
    ], widths=[1.15, 3.45, 2.4], font_size=9.5)

    add_heading(doc, "8 Kết luận tiến độ", 1)
    add_body(doc, "Đến thời điểm báo cáo, đồ án đã hoàn thành một sản phẩm demo có thể truy cập trực tuyến và sử dụng checkpoint UniXcoder thật. Hệ thống hỗ trợ phân tích mã C/C++, giải thích token, so sánh trước và sau bản vá, trình bày dashboard và tải kết quả. Pipeline huấn luyện, đánh giá, suy luận, đóng gói và triển khai đã được kiểm chứng.")
    add_body(doc, "Công việc trọng tâm tiếp theo không phải mở rộng giao diện mà là nâng chất lượng thực nghiệm: kiểm soát rò rỉ dữ liệu, xây dựng baseline, mở rộng fine-tune, chạy nhiều seed, thực hiện ablation và phân tích lỗi. Các kết quả hiện tại sẽ được giữ làm mốc baseline để so sánh với các phiên bản cải tiến.")

    # Footer page numbers
    for section in doc.sections:
        footer = section.footer
        add_page_number(footer.paragraphs[0])

    doc.core_properties.title = "Báo cáo tiến độ đồ án tốt nghiệp VulnScope"
    doc.core_properties.subject = "Ứng dụng UniXcoder trong phát hiện lỗi mã nguồn"
    doc.core_properties.author = "Nguyễn Văn Vinh"
    doc.core_properties.keywords = "UniXcoder, VulnScope, Devign, vulnerability detection"
    doc.save(OUTPUT)
    print(OUTPUT.name)


if __name__ == "__main__":
    build()
