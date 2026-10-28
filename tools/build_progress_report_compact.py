from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Bao_cao_tien_do_VulnScope_rut_gon.docx"


def font(run, size=12, bold=False):
    run.font.name = "Times New Roman"
    fonts = run._element.get_or_add_rPr().rFonts
    for name in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{name}"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def paragraph_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, line=1.3):
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line


def heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph_format(p, WD_ALIGN_PARAGRAPH.LEFT, before=10, after=6, line=1.15)
    for run in p.runs:
        font(run, 15 if level == 1 else 13, True)
    p.paragraph_format.keep_with_next = True
    return p


def body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    paragraph_format(p)
    if bold_lead and text.startswith(bold_lead):
        font(p.add_run(bold_lead), bold=True)
        font(p.add_run(text[len(bold_lead):]))
    else:
        font(p.add_run(text))
    return p


def bullet(doc, text):
    p = doc.add_paragraph()
    paragraph_format(p, after=3, line=1.2)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    font(p.add_run("• "))
    font(p.add_run(text))


def number(doc, text):
    p = doc.add_paragraph(style="List Number")
    paragraph_format(p, after=3, line=1.2)
    font(p.add_run(text))


def cell_margins(cell, value=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "8")
        node.set(qn("w:color"), "B7B7B7")
        borders.append(node)


def add_table(doc, headers, rows, widths, font_size=10.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    header_pr.append(OxmlElement("w:tblHeader"))
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell_margins(cell)
        p = cell.paragraphs[0]
        paragraph_format(p, WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.05)
        font(p.add_run(value), font_size, True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell_margins(cell)
            p = cell.paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.CENTER if index == 0 or len(str(value)) <= 12 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph_format(p, align, after=0, line=1.08)
            font(p.add_run(str(value)), font_size)
    table_borders(table)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table


def hyperlink(paragraph, text, url):
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([fonts, color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([props, text_node])
    link.append(run)
    paragraph._p.append(link)


def page_number(p):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])
    font(run, 10)


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.8)

    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    border = title_ppr.find(qn("w:pBdr"))
    if border is not None:
        title_ppr.remove(border)

    title = doc.add_paragraph(style="Title")
    paragraph_format(title, WD_ALIGN_PARAGRAPH.CENTER, after=14, line=1.1)
    font(title.add_run("BÁO CÁO TIẾN ĐỘ DEMO VULNSCOPE"), 18, True)

    heading(doc, "1 Khối lượng công việc đã hoàn thành", 1)
    add_table(doc, ["STT", "Hạng mục", "Kết quả", "Trạng thái"], [
        ("1", "Chuẩn bị dữ liệu", "Đọc được JSON và JSONL; chuẩn hóa các trường idx, func, target; chuẩn bị split Devign.", "Hoàn thành"),
        ("2", "Huấn luyện", "Fine-tune mô hình phân loại nhị phân, lưu checkpoint tốt nhất theo F1.", "Hoàn thành bản đầu"),
        ("3", "Suy luận", "Nạp checkpoint, dự đoán một mẫu hoặc theo batch, trả về xác suất và độ trễ.", "Hoàn thành"),
        ("4", "Giải thích", "Đánh giá ảnh hưởng token bằng phương pháp token occlusion.", "Hoàn thành"),
        ("5", "Đánh giá", "Tính Accuracy, Precision, Recall, F1, ROC-AUC, PR-AUC và confusion matrix.", "Hoàn thành bản đầu"),
        ("6", "Giao diện", "Xây dựng trang phân tích, so sánh bản vá và dashboard thực nghiệm.", "Hoàn thành"),
        ("7", "Triển khai", "Đưa mã nguồn lên GitHub, checkpoint lên Hugging Face và demo lên Streamlit Cloud.", "Hoàn thành"),
        ("8", "Kiểm thử", "Unit test dữ liệu, metric, bootstrap; smoke test Docker và checkpoint thật.", "Hoàn thành"),
    ], [0.45, 1.3, 4.0, 1.05])

    heading(doc, "3 Kiến trúc và quy trình xử lý", 1)
    body(doc, "VulnScope được tổ chức theo các mô-đun độc lập để tách quá trình chuẩn bị dữ liệu, huấn luyện, đánh giá và triển khai. Cách tổ chức này giúp kiểm thử từng thành phần và hạn chế việc giao diện can thiệp vào kết quả mô hình.")
    heading(doc, "3.1 Các thành phần chính", 2)
    add_table(doc, ["Thành phần", "Vai trò"], [
        ("data_utils.py", "Đọc và chuẩn hóa dữ liệu đầu vào từ JSON hoặc JSONL."),
        ("model.py", "Bao đóng encoder UniXcoder và lớp phân loại nhị phân."),
        ("run.py", "Huấn luyện, đánh giá theo epoch và lưu checkpoint tốt nhất."),
        ("inference.py", "Nạp checkpoint, dự đoán và giải thích token."),
        ("evaluate.py", "Đánh giá test set và tạo artifact cho dashboard."),
        ("app.py", "Giao diện Streamlit cho người dùng cuối."),
        ("bootstrap.py", "Chuẩn bị checkpoint và khởi động ứng dụng production."),
    ], [1.65, 5.15])
    heading(doc, "3.2 Luồng xử lý một yêu cầu", 2)
    for item in [
        "Người dùng chọn ví dụ, nhập trực tiếp một hàm hoặc tải file C/C++.",
        "Hệ thống chuẩn hóa khoảng trắng và tokenize bằng tokenizer của UniXcoder.",
        "Chuỗi token được thêm các token điều khiển, cắt theo block size và padding về cùng độ dài.",
        "Mô hình sinh xác suất cho hai lớp safe và vulnerable.",
        "Giao diện so sánh xác suất với ngưỡng cảnh báo và hiển thị kết quả.",
        "Khi người dùng yêu cầu giải thích, hệ thống lần lượt che token để đo thay đổi xác suất.",
    ]:
        number(doc, item)

    heading(doc, "4 Các chức năng của demo VulnScope", 1)
    heading(doc, "4.1 Phân tích mã nguồn", 2)
    body(doc, "Trang Phân tích tiếp nhận một hàm C/C++ từ ô nhập hoặc file có phần mở rộng c, cc, cpp, h và hpp. Người dùng cũng có thể chọn các ví dụ buffer overflow, use-after-free hoặc phiên bản an toàn. Sau khi phân tích, hệ thống hiển thị xác suất lỗ hổng, độ tin cậy, số token, độ trễ và cảnh báo khi đầu vào bị cắt.")
    for item in [
        "Kết quả được lưu trong session để không biến mất khi giao diện chạy lại.",
        "Người dùng có thể tải kết quả ở định dạng JSON.",
        "Production bắt buộc dùng checkpoint thật và không tự chuyển sang chế độ minh họa.",
        "Kích thước file tải lên được giới hạn để giảm rủi ro tài nguyên.",
    ]:
        bullet(doc, item)
    heading(doc, "4.2 Giải thích token", 2)
    body(doc, "Chức năng giải thích sử dụng token occlusion. Hệ thống dự đoán xác suất ban đầu, sau đó lần lượt thay từng token bằng mask token và dự đoán lại. Chênh lệch xác suất được dùng làm attribution. Attribution dương cho biết token làm tăng dự đoán lỗ hổng; giá trị tuyệt đối lớn cho biết token có ảnh hưởng mạnh.")
    heading(doc, "4.3 So sánh trước và sau bản vá", 2)
    body(doc, "Người dùng nhập hai phiên bản của cùng một hàm. Hệ thống sử dụng cùng checkpoint, block size và ngưỡng để đánh giá hai phiên bản, sau đó hiển thị xác suất trước vá, sau vá, mức thay đổi và phần diff của mã nguồn. Chức năng này hỗ trợ quan sát liệu bản vá có làm giảm rủi ro theo mô hình hay không.")
    heading(doc, "4.4 Dashboard thực nghiệm", 2)
    body(doc, "Dashboard chỉ đọc artifact do chương trình đánh giá tạo ra. Các chỉ số không được sinh giả trong giao diện. Trang hiện hiển thị Accuracy, Precision, Recall, F1, ROC-AUC, PR-AUC, confusion matrix và biểu đồ phân bố true negative, false positive, false negative, true positive.")

    heading(doc, "5 Triển khai và kiểm thử", 1)
    heading(doc, "5.1 Triển khai trực tuyến", 2)
    p = doc.add_paragraph()
    paragraph_format(p, WD_ALIGN_PARAGRAPH.LEFT)
    font(p.add_run("Mã nguồn được quản lý trong một repository độc lập trên GitHub. Checkpoint dung lượng khoảng 480 MB được lưu trong Hugging Face Model Repository thay vì commit trực tiếp vào Git. Ứng dụng được triển khai trên Streamlit Community Cloud và tự tải checkpoint đã ghim phiên bản khi khởi động."))
    for label, text, url in [
        ("Demo trực tuyến: ", "https://vulnscope-unixcoder.streamlit.app", "https://vulnscope-unixcoder.streamlit.app"),
        ("Mã nguồn: ", "https://github.com/vinh204/vulnscope-unixcoder", "https://github.com/vinh204/vulnscope-unixcoder"),
        ("Checkpoint: ", "https://huggingface.co/vinh204/vulnscope-unixcoder-devign", "https://huggingface.co/vinh204/vulnscope-unixcoder-devign"),
    ]:
        p = doc.add_paragraph()
        paragraph_format(p, WD_ALIGN_PARAGRAPH.LEFT, after=5, line=1.15)
        font(p.add_run(label), bold=True)
        font(p.add_run(text))
    heading(doc, "5.2 Các biện pháp ổn định và an toàn", 2)
    for item in [
        "Đóng gói bằng Docker và chạy container bằng user không có quyền root.",
        "Cung cấp health check để nền tảng giám sát trạng thái ứng dụng.",
        "Dùng cache để chỉ nạp tokenizer và model một lần trong mỗi tiến trình.",
        "Tách checkpoint và dataset lớn khỏi Git repository bằng gitignore và dockerignore.",
        "Không thực thi mã nguồn do người dùng tải lên; hệ thống chỉ tokenize và suy luận.",
        "Ghim revision của checkpoint để deployment có khả năng tái lập.",
    ]:
        bullet(doc, item)
    heading(doc, "5.3 Kết quả kiểm thử", 2)
    add_table(doc, ["Nhóm kiểm thử", "Nội dung", "Kết quả"], [
        ("Unit test", "Đọc JSON, JSONL; metric khi đủ lớp và thiếu lớp; bootstrap checkpoint.", "6 trên 6 đạt"),
        ("Syntax", "Biên dịch kiểm tra toàn bộ thư mục code.", "Đạt"),
        ("Docker", "Build image, khởi động container, kiểm tra health endpoint.", "Đạt"),
        ("Inference", "Nạp checkpoint thật và chạy dự đoán mẫu buffer overflow.", "Đạt"),
        ("Production", "URL public trả HTTP 200 và giao diện hoạt động.", "Đạt"),
    ], [1.3, 4.55, 1.0])

    heading(doc, "6 Kết quả thực nghiệm ban đầu", 1)
    body(doc, "Checkpoint hiện tại được xây dựng để xác minh toàn bộ pipeline và phục vụ demo. Encoder được đóng băng, mô hình được huấn luyện 3 epoch trên 2.000 mẫu, block size 128 và đánh giá trên 500 mẫu test. Ngưỡng cảnh báo đang sử dụng là 0,44.")
    add_table(doc, ["Chỉ số", "Giá trị", "Ý nghĩa"], [
        ("Accuracy", "0,504", "Tỷ lệ dự đoán đúng trên toàn bộ mẫu."),
        ("Precision", "0,486", "Tỷ lệ cảnh báo vulnerable thực sự là vulnerable."),
        ("Recall", "0,579", "Tỷ lệ mẫu vulnerable được phát hiện."),
        ("F1", "0,529", "Trung bình điều hòa giữa precision và recall."),
        ("ROC-AUC", "0,528", "Khả năng phân biệt hai lớp trên nhiều ngưỡng."),
        ("PR-AUC", "0,529", "Quan hệ precision-recall trên nhiều ngưỡng."),
    ], [1.35, 1.0, 4.5])
    heading(doc, "6.1 Ma trận nhầm lẫn", 2)
    add_table(doc, ["Nhãn thật", "Dự đoán an toàn", "Dự đoán vulnerable"], [
        ("An toàn", "113", "147"),
        ("Vulnerable", "101", "139"),
    ], [2.2, 2.2, 2.45])
    body(doc, "Nhận xét: kết quả hiện tại gần mức baseline và chưa đủ để kết luận UniXcoder đạt hiệu quả tốt trên Devign. Kết quả này phản ánh đúng cấu hình thử nghiệm hạn chế: tập huấn luyện nhỏ, encoder đóng băng và block size ngắn. Giá trị của giai đoạn hiện tại là chứng minh pipeline hoạt động từ đầu đến cuối; kết quả cuối cần được xác lập sau khi mở rộng huấn luyện và thực hiện đối chứng.", "Nhận xét:")

    page_number(doc.sections[0].footer.paragraphs[0])
    doc.core_properties.title = "Báo cáo tiến độ demo VulnScope"
    doc.core_properties.author = "Nguyễn Văn Vinh"
    doc.save(OUTPUT)
    print(OUTPUT.name)


if __name__ == "__main__":
    build()
