from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "Chuong_3.md"
OUTPUT = ROOT / "docs" / "Chuong_3_UniXcoder.docx"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    properties.append(shade)


def set_cell_text(cell, text, header=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text.strip())
    run.bold = header
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    if header:
        run.font.color.rgb = RGBColor(255, 255, 255)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, value, end):
        run._r.append(node)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)


def parse_table(lines, document):
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
            continue
        rows.append(cells)
    if not rows:
        return
    columns = max(map(len, rows))
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row in enumerate(rows):
        for column_index in range(columns):
            value = row[column_index] if column_index < len(row) else ""
            cell = table.cell(row_index, column_index)
            set_cell_text(cell, value, header=row_index == 0)
            if row_index == 0:
                set_cell_shading(cell, "244A67")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "EEF3F7")
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_body_paragraph(document, text, style=None):
    paragraph = document.add_paragraph(style=style)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.28)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.add_run(text)
    return paragraph


def main():
    lines = SOURCE.read_text(encoding="utf-8-sig").splitlines()
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.15)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.line_spacing = 1.5
    for name, size, color in (
        ("Heading 1", 15, "18344E"),
        ("Heading 2", 13, "244A67"),
        ("Heading 3", 12, "365F7D"),
    ):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "VULNSCOPE — PHÁT HIỆN LỖ HỔNG MÃ NGUỒN"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(100, 110, 120)
    add_page_number(section.footer.paragraphs[0])

    index = 0
    while index < len(lines):
        raw = lines[index].strip()
        if not raw:
            index += 1
            continue
        if raw.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            parse_table(table_lines, document)
            continue
        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", raw)
        if image_match:
            image_path = ROOT / "docs" / image_match.group(2)
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.add_run().add_picture(str(image_path), width=Inches(6.0))
            index += 1
            continue
        if raw.startswith("**") and raw.endswith("**"):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(8)
            run = paragraph.add_run(raw.strip("*"))
            run.italic = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(10)
            index += 1
            continue
        if raw.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(14)
            run = paragraph.add_run(raw[2:])
            run.font.name = "Times New Roman"
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(24, 52, 78)
            index += 1
            continue
        if raw.startswith("### "):
            document.add_heading(raw[4:], level=3)
            index += 1
            continue
        if raw.startswith("## "):
            document.add_heading(raw[3:], level=2)
            index += 1
            continue
        if raw.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.3)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.add_run(raw[2:])
            index += 1
            continue
        add_body_paragraph(document, raw)
        index += 1

    document.core_properties.title = "Chương 3 — Ứng dụng UniXcoder trong phát hiện lỗi mã nguồn"
    document.core_properties.subject = "Đồ án VulnScope"
    document.core_properties.author = "VulnScope"
    document.save(OUTPUT)
    print("DOCX generated successfully.")
    print(f"Markdown word count: {len(SOURCE.read_text(encoding='utf-8-sig').split())}")


if __name__ == "__main__":
    main()
